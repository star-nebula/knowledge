#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 xlsx 的指定 sheet 可靠导出为 CSV（UTF-8 BOM）。

用法: python extract_sheet.py <xlsx> [sheet名] [out_dir]
默认 sheet 名: 8.3~
⚠️ 必须用普通 load_workbook（非 read_only），否则此文件会被误读成 1 行。
"""
import sys
import os
import csv
import openpyxl


def main():
    if len(sys.argv) < 2:
        print("用法: extract_sheet.py <xlsx> [sheet名] [out_dir]")
        sys.exit(1)
    xlsx = sys.argv[1]
    sheet = sys.argv[2] if len(sys.argv) > 2 else "8.3~"
    out_dir = sys.argv[3] if len(sys.argv) > 3 else os.path.dirname(xlsx)
    os.makedirs(out_dir, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx, read_only=False, data_only=True)  # 必须非 read_only
    if sheet not in wb.sheetnames:
        print("可用 sheet:", wb.sheetnames)
        sys.exit(1)
    ws = wb[sheet]
    max_r, max_c = ws.max_row, ws.max_column
    out = os.path.join(out_dir, sheet + ".csv")
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        n = 0
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row[:max_c]])
            n += 1
    print(f"导出 {sheet}: {n} 行 x {max_c} 列 -> {out}")


if __name__ == "__main__":
    main()
