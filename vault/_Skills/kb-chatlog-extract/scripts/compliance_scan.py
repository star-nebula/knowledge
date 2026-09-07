#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""合规扫描卡点：扫描批次 CSV 的 提炼后_问题 / 提炼后_答案 列，标记疾病宣称等违规词。

用法: python compliance_scan.py <批次.csv>
- 仅做关键词标记，是否真违规由 LLM 判断并改写；命中行需修正后才可 入库。
- 自动处理 CSV 可能被转成 zip（PK 头）：尝试用 openpyxl 恢复读取。
- 疾病宣称词表针对家具/枕头/按摩类"不得宣称治疗疾病"红线（治疗/治愈/根治/理疗/颈椎病/腰椎/
  腰疼/腰痛/失眠/关节炎…）；舒缓/辅助/缓解 不算违规，不列入。
"""
import sys
import os
import io
import csv
import zipfile
import openpyxl

# 疾病宣称词（命中即需改写为舒缓/辅助表述）；注意不含 舒缓/辅助/缓解
DISEASE_WORDS = [
    "治疗", "治愈", "根治", "理疗", "颈椎病", "腰椎", "腰间盘突出", "腰疼", "腰痛",
    "腰酸痛", "失眠", "关节炎", "坐骨神经", "病症", "疾病", "活血", "化瘀", "镇痛",
]


def load_rows(path):
    with open(path, "rb") as f:
        head = f.read(2)
    if head == b"PK":
        with open(path, "rb") as f:
            blob = f.read()
        wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=False, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        wb.close()
        hdr = rows[0]
        return [dict(zip(hdr, r)) for r in rows[1:]]
    with io.open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def main():
    if len(sys.argv) < 2:
        print("用法: compliance_scan.py <批次.csv>")
        sys.exit(1)
    p = sys.argv[1]
    rows = load_rows(p)
    hits = 0
    for r in rows:
        q = r.get("提炼后_问题") or ""
        a = r.get("提炼后_答案") or ""
        text = q + " " + a
        found = [w for w in DISEASE_WORDS if w in text]
        if found:
            hits += 1
            print(f"[命中] 会话={r.get('会话ID')} | 结论={r.get('入库结论')} | 词={','.join(found)}")
            print(f"       问题: {q[:60]}")
            print(f"       答案: {a[:90]}")
    print(f"\n扫描完成: {len(rows)} 行, 命中 {hits} 行待复核改写。")
    if hits:
        print("→ 请对命中行去掉疾病宣称（改舒缓/辅助表述）后再 入库。")


if __name__ == "__main__":
    main()
