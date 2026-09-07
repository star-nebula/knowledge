#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 LLM 产出的候选条目 JSON 安全写出为批次 CSV（UTF-8 BOM），并防 zip 损坏。

用法: python write_batch.py <rows.json>

JSON 结构:
{
  "batch": 4,
  "out": "可选覆盖路径，缺省为同目录 知识提炼_批次{batch}.csv",
  "rows": [ {"批次","会话ID","商品","提炼前_原始问答","提炼后_问题","提炼后_答案",
             "意图","答案标签","适用对象","优先级","入库结论","备注"}, ... ]
}

- 写出后校验文件头非 PK；若被转成 zip，自动用 openpyxl 从 zip 恢复（read_only=False 拿全 12 列）并重写。
- 打印 入库/复核/不入库 计数。
"""
import sys
import os
import csv
import json
import shutil
import zipfile
import openpyxl

COLS = ["批次", "会话ID", "商品", "提炼前_原始问答", "提炼后_问题", "提炼后_答案",
        "意图", "答案标签", "适用对象", "优先级", "入库结论", "备注"]


def recover_from_zip(path):
    """文件被转成 zip：用 openpyxl 读回全 12 列并重写为 CSV。

    用 BytesIO 直接读 zip 字节，不落盘临时 .xlsx（避开沙箱 safe-delete 限制）。
    """
    import io
    with open(path, "rb") as f:
        blob = f.read()
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    data = [[c.value for c in r] for r in ws.iter_rows()]
    wb.close()
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerows(data)


def main():
    if len(sys.argv) < 2:
        print("用法: write_batch.py <rows.json>")
        sys.exit(1)
    jp = sys.argv[1]
    with open(jp, encoding="utf-8-sig") as f:
        data = json.load(f)
    rows = data["rows"]
    out = data.get("out")
    if not out:
        d = os.path.dirname(jp) or "."
        bn = data.get("batch") or 1
        out = os.path.join(d, f"知识提炼_批次{bn}.csv")

    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(rows)

    # 校验是否被转成 zip
    with open(out, "rb") as f:
        head = f.read(4)
    if head[:2] == b"PK":
        print("⚠️ 写出后文件被转成 zip，正在恢复...")
        recover_from_zip(out)

    def _cat(c):
        c = (c or "").strip()
        if "复核" in c:
            return "复核"
        if c == "入库":
            return "入库"
        if c == "不入库":
            return "不入库"
        return "其他"

    ru = sum(1 for r in rows if _cat(r.get("入库结论")) == "入库")
    rc = sum(1 for r in rows if _cat(r.get("入库结论")) == "复核")
    ni = sum(1 for r in rows if _cat(r.get("入库结论")) == "不入库")
    print(f"已写出: {out} | 行数 {len(rows)} (入库{ru}/复核{rc}/不入库{ni}) | 大小 {os.path.getsize(out)}")

    with open(out, "rb") as f:
        assert f.read(2) != b"PK", "仍为 zip，恢复失败！"
    print("✅ 文件头校验通过（非 zip）")


if __name__ == "__main__":
    main()
