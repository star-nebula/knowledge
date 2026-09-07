#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""从 8.3~.csv 选出下一批待提炼会话，并导出原始对话子集 CSV 供 LLM 直接阅读。

用法: python select_next_batch.py <8.3~.csv> <输出目录> [--count 5]

- 自动跳过已做会话：扫描 输出目录/知识提炼_批次*.csv 收集 done 会话ID。
- 选批规则：3 个带商品（尽量不同产品）+ 余下用无商品会话补足到 count，且都有实质问答。
- 把选中会话全部原始消息写到 输出目录/批次N_原始对话.csv（保留完整 9 列、不截断、不重排），LLM 用 Read 工具直接读该子集 CSV。
- 实质问答判定：客户真实提问 + 客服非套话/非占位答复（跳过欢迎语/占位"请在客户端查看"）。
"""
import sys
import os
import csv
import io
import zipfile
import openpyxl
from collections import defaultdict

PLACE = "请在客户端查看原始聊天记录"
GREET = {"你好", "在吗", "您好", "嗯", "哦", "好的", "谢谢", "在的", "亲亲"}
BOILER = ["Hi~我们是致力于", "下单前咱们确认好收货信息", "亲，收货的时候麻烦多留意"]


def is_boiler(t):
    t = (t or "").strip()
    return (not t) or any(t.startswith(p) for p in BOILER) or t in (
        "在的呢亲亲", "ovo 在的，亲亲，请问有什么可以帮您的呢？",
        "您可以看看的哦~", "嗯呢亲亲~")


def is_real_q(t):
    t = (t or "").strip()
    return t not in GREET and PLACE not in t


def real_qa(msgs):
    n = len(msgs)
    c = 0
    i = 0
    while i < n:
        m = msgs[i]
        if m["对象"] == "客户" and is_real_q(m["消息"]):
            j = i + 1
            got = False
            while j < n:
                cj = msgs[j]
                if cj["对象"] == "客户":
                    if is_real_q(cj["消息"]):
                        break
                    else:
                        j += 1
                        continue
                if PLACE in (cj["消息"] or "") or is_boiler(cj["消息"]):
                    j += 1
                    continue
                got = True
                j += 1
            if got:
                c += 1
            i = j if j > n else j
        else:
            i += 1
    return c


def has_prod(msgs):
    return any((m["商品ID"].strip() or m["商品"].strip()) for m in msgs)


def recover_csv_from_zip(path):
    """CSV 被环境转成 zip(PK)：用 openpyxl 读回全部列并重写为 CSV（不落盘临时文件）。"""
    with open(path, "rb") as f:
        blob = f.read()
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    wb.close()
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerows(rows)


def collect_done(out_dir):
    done = set()
    for fn in os.listdir(out_dir):
        # 只认规范交付文件 知识提炼_批次N.csv；排除 _agent / _合并 / 待审 等中间文件
        if (fn.startswith("知识提炼_批次") and fn.endswith(".csv")
                and "_agent" not in fn and "_合并" not in fn and "待审" not in fn):
            fp = os.path.join(out_dir, fn)
            try:
                # 自愈：若被转成 zip，先从 zip 恢复为 CSV 再读
                if zipfile.is_zipfile(fp):
                    with open(fp, "rb") as f:
                        blob = f.read()
                    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=False, data_only=True)
                    ws = wb[wb.sheetnames[0]]
                    rows = [[c.value for c in r] for r in ws.iter_rows()]
                    wb.close()
                    with open(fp, "w", encoding="utf-8-sig", newline="") as f:
                        w = csv.writer(f)
                        w.writerows(rows)
                with open(fp, encoding="utf-8-sig", newline="") as f:
                    for row in csv.DictReader(f):
                        if row.get("会话ID"):
                            done.add(row["会话ID"])
            except Exception:
                pass
    return done


def next_batch_no(out_dir):
    bn = 1
    for fn in os.listdir(out_dir):
        if (fn.startswith("知识提炼_批次") and fn.endswith(".csv")
                and "_agent" not in fn and "_合并" not in fn and "待审" not in fn):
            try:
                bn = max(bn, int(fn.split("批次")[1].split(".")[0]) + 1)
            except Exception:
                pass
    return bn


def main():
    if len(sys.argv) < 3:
        print("用法: select_next_batch.py <8.3~.csv> <输出目录> [--count 5]")
        sys.exit(1)
    csv_path = sys.argv[1]
    out_dir = sys.argv[2]
    count = 5
    if "--count" in sys.argv:
        count = int(sys.argv[sys.argv.index("--count") + 1])
    os.makedirs(out_dir, exist_ok=True)

    done = collect_done(out_dir)
    rows = []
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            rows.append(r)
    sessions = defaultdict(list)
    for r in rows:
        sessions[r["会话ID"]].append(r)

    withp, without = [], []
    for sid, msgs in sessions.items():
        if sid in done:
            continue
        c = real_qa(msgs)
        if c <= 0:
            continue
        if has_prod(msgs):
            withp.append((sid, c, msgs[0]["商品"]))
        else:
            without.append((sid, c))

    withp.sort(key=lambda x: -x[1])
    picked, seenp = [], set()
    for sid, c, prod in withp:
        if len(picked) >= 3:  # 带商品最多 3 个
            break
        if prod in seenp:
            continue
        picked.append(sid)
        seenp.add(prod)
    for sid, c in without:
        if len(picked) >= count:
            break
        picked.append(sid)

    batch_no = next_batch_no(out_dir)
    print("已做会话数:", len(done))
    print("选中会话:", picked, "| 批次号:", batch_no)

    dump = os.path.join(out_dir, f"批次{batch_no}_原始对话.csv")
    cols = ["会话ID", "日期", "时间", "对象", "发话者", "消息", "消息状态", "商品ID", "商品"]
    with open(dump, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for sid in picked:
            for m in sessions[sid]:
                w.writerow({c: (m.get(c) or "") for c in cols})
    print("原始对话(子集 CSV)已写入:", dump)
    print("用 Read 工具读取该 CSV，按 SKILL.md 提炼规则产出候选条目，再用 write_batch.py 写出。")

    # 防损坏：环境可能把 CSV 转成 zip(PK)，写出后立即校验并自愈
    with open(dump, "rb") as f:
        if f.read(2) == b"PK":
            print("⚠️ 原始对话 CSV 被转成 zip，正在自愈恢复...")
            recover_csv_from_zip(dump)


if __name__ == "__main__":
    main()
